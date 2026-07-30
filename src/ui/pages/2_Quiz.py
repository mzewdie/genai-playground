import io
import json
import logging

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from src.common.llm.llm_configuration import LLMConfiguration
from src.common.llm.llm_factory import LLMFactory
from src.common.llm.llm_service import LLMService
from src.document_analyzer.action import Action
from src.document_analyzer.document_analysis_orchestrator import DocumentAnalysisOrchestrator
from src.document_analyzer.quiz_service import QuizService

logger = logging.getLogger(__name__)

st.title("Quiz Page")

if "selected_document" in st.session_state:
    st.write(f"Selected Document: {st.session_state['selected_document']}")
else:
    st.error("You must first load the document from the Sidebar Load Document")
    st.stop()

if "selected_llm_provider" in st.session_state:
    st.write(f"LLM: {st.session_state['selected_llm_provider']}")
if "selected_llm_model" in st.session_state:
    st.write(f"Model: {st.session_state['selected_llm_model']}")

quiz_button = st.button("Generate Quiz")

###############################################################################
def format_options_list(options: list) -> list[str]:
    """
    Transforms a list of options ['Apple', 'Banana'] into ['A. Apple', 'B. Banana'].
    """
    if not options:
        return []
    return [f"{chr(65 + i)}. {opt}" for i, opt in enumerate(options)]


def get_formatted_correct_answer(question: dict) -> str:
    """
    Returns a unified string representation of the correct answer.
    """
    q_type = question.get("type")
    correct_val = question.get("correct_answer")

    if q_type == "multiple_choice":
        options = question.get("options", [])
        if isinstance(correct_val, int) and 0 <= correct_val < len(options):
            return f"{chr(65 + correct_val)}. {options[correct_val]}"
        return str(correct_val)

    elif q_type == "true_false":
        return "True" if str(correct_val).lower() in ["true", "1"] else "False"

    return str(correct_val)



##############################################
# -------------------------------------------------------------
# 1. Pure-Python PDF Generator (Runs on Server)
# -------------------------------------------------------------
def build_pdf_stream(quiz_data: dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle", parent=styles["Heading1"], fontSize=18, textColor=colors.HexColor("#1A365D")
    )
    q_style = ParagraphStyle(
        "QTitle", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#2B6CB0"), spaceBefore=10
    )
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=14)
    answer_style = ParagraphStyle(
        "Answer",
        parent=styles["Normal"],
        fontSize=9.5,
        textColor=colors.HexColor("#22543D"),
        backColor=colors.HexColor("#F0FFF4"),
        borderColor=colors.HexColor("#38A169"),
        borderWidth=1,
        borderPadding=6,
        spaceBefore=6,
        spaceAfter=10,
    )

    story = [
        Paragraph("Quiz from the Document & Study Guide", title_style),
        Spacer(1, 4),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2B6CB0"), spaceAfter=15),
    ]

    for idx, q in enumerate(quiz_data.get("questions", []), start=1):
        q_type = q.get("type", "").replace("_", " ").title()
        story.append(Paragraph(f"Question {idx} ({q_type}):", q_style))
        story.append(Paragraph(q["question"], body_style))

        if q.get("type") == "multiple_choice":
            for o_idx, opt in enumerate(q["options"]):
                story.append(
                    Paragraph(f"<b>{chr(65 + o_idx)}.</b> {opt}", ParagraphStyle("Opt", parent=body_style, leftIndent=15))
                )
            correct_ans = f"<b>{chr(65 + q['correct_answer'])}.</b> {q['options'][q['correct_answer']]}"
        elif q.get("type") == "true_false":
            correct_ans = "True" if q["correct_answer"] else "False"
        else:
            correct_ans = str(q["correct_answer"])

        story.append(
            Paragraph(f"<b>Correct Answer:</b> {correct_ans}<br/><b>Explanation:</b> {q['explanation']}", answer_style)
        )
        story.append(Spacer(1, 6))

    doc.build(story)
    return buffer.getvalue()



###########################################################################################

def build_excel_stream(quiz_data: dict) -> bytes:
    """Generates a styled, non-empty Excel workbook from quiz_data."""
    # 1. Create a new OpenPyXL Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz Summary"

    # Ensure grid lines are visible in Excel
    ws.views.sheetView[0].showGridLines = True

    # 2. Style Definitions
    header_fill = PatternFill(
        start_color="1A365D", end_color="1A365D", fill_type="solid"
    )  # Dark Blue
    zebra_fill = PatternFill(
        start_color="F7FAFC", end_color="F7FAFC", fill_type="solid"
    )  # Light Gray
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10, color="000000")

    thin_border_side = Side(style="thin", color="D1D5DB")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    align_center_top = Alignment(
        horizontal="center", vertical="top", wrap_text=True
    )
    align_left_top = Alignment(
        horizontal="left", vertical="top", wrap_text=True
    )

    # 3. Write Headers
    headers = [
        "Q#",
        "Question Type",
        "Question",
        "Options",
        "Correct Answer",
        "Explanation",
    ]
    ws.append(headers)

    # Apply Header Styles
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center_top
        cell.border = thin_border

    # 4. Write Data Rows directly to worksheet
    questions = quiz_data.get("questions", [])

    for idx, q in enumerate(questions, start=1):
        q_type = q.get("type", "")

        # Format options
        if q_type == "multiple_choice":
            opts_list = q.get("options", [])
            opts_text = "\n".join(
                [f"{chr(65 + i)}. {opt}" for i, opt in enumerate(opts_list)]
            )

            correct_val = q.get("correct_answer")
            if isinstance(correct_val, int) and 0 <= correct_val < len(
                opts_list
            ):
                correct_text = (
                    f"{chr(65 + correct_val)}. {opts_list[correct_val]}"
                )
            else:
                correct_text = str(correct_val)
        elif q_type == "true_false":
            opts_text = "True\nFalse"
            correct_text = "True" if q.get("correct_answer") else "False"
        else:
            opts_text = "N/A"
            correct_text = str(q.get("correct_answer", ""))

        row_data = [
            idx,
            q_type.replace("_", " ").title(),
            q.get("question", ""),
            opts_text,
            correct_text,
            q.get("explanation", ""),
        ]

        ws.append(row_data)

        row_num = idx + 1  # Headers are on Row 1
        is_even = row_num % 2 == 0

        # Apply Row Styles
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = body_font
            cell.border = thin_border

            if is_even:
                cell.fill = zebra_fill

            # Alignment
            if col_num in [1, 2]:  # Q# and Type
                cell.alignment = align_center_top
            else:  # Text content
                cell.alignment = align_left_top

    # 5. Set Column Widths
    col_widths = {
        "A": 8,  # Q#
        "B": 18,  # Question Type
        "C": 45,  # Question
        "D": 30,  # Options
        "E": 25,  # Correct Answer
        "F": 45,  # Explanation
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 6. Save to BytesIO Memory Stream
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)  # Rewind buffer pointer to start

    return buffer.getvalue()


###########################################################################################
# -------------------------------------------------------------
# 2. Pure-Python Excel Generator (Runs on Server)
# -------------------------------------------------------------
def build_excel_stream_row(quiz_data: dict) -> bytes:
    rows = []
    for idx, q in enumerate(quiz_data.get("questions", []), start=1):
        q_type = q.get("type", "")
        if q_type == "multiple_choice":
            opts = " | ".join([f"{chr(65+i)}: {o}" for i, o in enumerate(q["options"])])
            correct = f"{chr(65+q['correct_answer'])}: {q['options'][q['correct_answer']]}"
        elif q_type == "true_false":
            opts = "True | False"
            correct = "True" if q["correct_answer"] else "False"
        else:
            opts = "N/A"
            correct = str(q["correct_answer"])

        rows.append({
            "Q#": idx,
            "Type": q_type.replace("_", " ").title(),
            "Question": q["question"],
            "Options": opts,
            "Correct Answer": correct,
            "Explanation": q["explanation"],
        })

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, index=False, sheet_name="Quiz")
    return buffer.getvalue()


# -------------------------------------------------------------
# 3. Streamlit Interface (What Users Interact With)
# -------------------------------------------------------------
#st.title("📄 Quiz Generator & Exporter")

if "quiz_json_data" in st.session_state and st.session_state["quiz_json_data"]:
    quiz_data = st.session_state["quiz_json_data"]

    st.subheader("Export Options")
    col1, col2 = st.columns(2)

    with col1:
        st.download_button(
            label="📄 Download Printable PDF",
            data=build_pdf_stream(quiz_data),
            file_name="Quiz_Study_Guide.pdf",
            mime="application/pdf",
        )

    with col2:
        st.download_button(
            label="📊 Download Excel Spreadsheet",
            data=build_excel_stream(quiz_data),
            file_name="Quiz_Data.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


def render_quiz_question(q: dict, idx: int):
    st.markdown(f"**Question {idx}:** {q['question']}")
    q_type = q.get("type")

    if q_type == "multiple_choice":
        # Apply unified option formatting
        formatted_options = format_options_list(q.get("options", []))

        st.radio(
            "Select your answer:",
            options=formatted_options,
            key=f"ui_q_{idx}",
            index=None
        )

    elif q_type == "true_false":
        st.radio(
            "Select True or False:",
            options=["True", "False"],
            key=f"ui_tf_{idx}",
            index=None
        )

    # Explanation expander
    with st.expander("Reveal Answer & Explanation"):
        correct_text = get_formatted_correct_answer(q)
        st.success(f"**Correct Answer:** {correct_text}")
        st.info(f"**Explanation:** {q.get('explanation', '')}")



###################

import json
import streamlit as st


# ==========================================
# 1. HELPER: Unified Option Formatting
# ==========================================
def format_options(options: list) -> list[str]:
    """Formats options into 'A. Option Text'."""
    return [f"{chr(65 + i)}. {opt}" for i, opt in enumerate(options)]


# ==========================================
# 2. MAIN QUIZ APP WITH SCORING STATE
# ==========================================
def render_scorable_quiz():


    # ######### Initialize quiz data in session state if missing
    #st.title("📝 Interactive Assessment")
    # if "quiz_json_data" not in st.session_state:
    #     sample_json = """{
    #         "questions": [
    #             {
    #                 "type": "multiple_choice",
    #                 "question": "What is the primary function of carbohydrates in human nutrition?",
    #                 "options": ["Building muscle tissue", "Providing immediate energy", "Storing genetic information"],
    #                 "correct_answer": 1,
    #                 "explanation": "Carbohydrates serve as the primary and fastest source of energy for the body."
    #             },
    #             {
    #                 "type": "true_false",
    #                 "question": "Water is considered one of the six essential nutrient classes.",
    #                 "correct_answer": true,
    #                 "explanation": "Water is essential for cellular function and hydration."
    #             }
    #         ]
    #     }"""
    #     st.session_state["quiz_json_data"] = json.loads(sample_json)
    # ####################################

    # Initialize user state containers
    if "user_answers" not in st.session_state:
        st.session_state["user_answers"] = {}
    if "quiz_submitted" not in st.session_state:
        st.session_state["quiz_submitted"] = False

    questions = st.session_state["quiz_json_data"].get("questions", [])

    # Create a Form to group all inputs together
    with st.form(key="quiz_form"):
        st.subheader("Questions")

        for idx, q in enumerate(questions, start=1):
            st.markdown(f"**Question {idx}:** {q['question']}")
            q_type = q.get("type")
            widget_key = f"user_q_{idx}"

            if q_type == "multiple_choice":
                opts = format_options(q.get("options", []))
                st.radio(
                    "Select your answer:",
                    options=opts,
                    key=widget_key,
                    index=None,
                    disabled=st.session_state["quiz_submitted"],
                )

            elif q_type == "true_false":
                st.radio(
                    "Select True or False:",
                    options=["True", "False"],
                    key=widget_key,
                    index=None,
                    disabled=st.session_state["quiz_submitted"],
                )

            st.divider()

        # Submit button inside the form
        submit_button = st.form_submit_button(
            label="Submit Quiz for Grading",
            disabled=st.session_state["quiz_submitted"],
        )

    # Process score upon form submission
    if submit_button:
        st.session_state["quiz_submitted"] = True
        st.rerun()

    # ==========================================
    # 3. SCORE CALCULATION & RESULTS DISPLAY
    # ==========================================
    if st.session_state["quiz_submitted"]:
        total_questions = len(questions)
        correct_count = 0

        st.subheader("📊 Your Quiz Results")

        for idx, q in enumerate(questions, start=1):
            user_choice = st.session_state.get(f"user_q_{idx}")
            q_type = q.get("type")
            is_correct = False

            # Check correctness by type
            if q_type == "multiple_choice":
                correct_idx = q.get("correct_answer")
                expected_prefix = f"{chr(65 + correct_idx)}."
                if user_choice and user_choice.startswith(expected_prefix):
                    is_correct = True
                    correct_text = (
                        f"{chr(65 + correct_idx)}. {q['options'][correct_idx]}"
                    )
                else:
                    correct_text = (
                        f"{chr(65 + correct_idx)}. {q['options'][correct_idx]}"
                    )

            elif q_type == "true_false":
                expected_str = "True" if q.get("correct_answer") else "False"
                if user_choice == expected_str:
                    is_correct = True
                correct_text = expected_str

            if is_correct:
                correct_count += 1
                st.success(
                    f"**Question {idx}: Correct!** You selected: `{user_choice}`"
                )
            else:
                selected_display = (
                    f"`{user_choice}`" if user_choice else "*Unanswered*"
                )
                st.error(
                    f"**Question {idx}: Incorrect.** You selected: {selected_display}"
                )
                st.info(
                    f"**Correct Answer:** {correct_text}\n\n**Explanation:** {q.get('explanation')}"
                )

        # Final Score Summary Box
        score_pct = int((correct_count / total_questions) * 100)

        st.markdown("---")
        col1, col2 = st.columns(2)
        with col1:
            st.metric(
                label="Final Score",
                value=f"{correct_count} / {total_questions}",
            )
        with col2:
            st.metric(label="Percentage", value=f"{score_pct}%")

        if score_pct >= 70:
            st.balloons()
            st.success("🎉 Great job! You passed the quiz.")
        else:
            st.warning("📚 Keep studying and try again!")

        # Reset button to allow retaking the quiz
        if st.button("🔄 Retake Quiz"):
            st.session_state["quiz_submitted"] = False
            # Clear individual answer keys
            for idx in range(1, total_questions + 1):
                if f"user_q_{idx}" in st.session_state:
                    del st.session_state[f"user_q_{idx}"]
            st.rerun()





########################

#json_test_str="""{ "questions": [ { "type": "multiple_choice", "question": "What is the definition of Nutrition according to the document?", "options": [ "The process of preparing food for consumption.", "The science of food and its relation to health.", "The study of cooking techniques and food preservation.", "The practice of consuming only organic foods." ], "correct_answer": 1, "explanation": "The document defines nutrition as 'the science of food and its relation to health.'" }, { "type": "true_false", "question": "The human body is capable of producing all 22 different amino acids required for protein synthesis.", "correct_answer": false, "explanation": "The document states that the human body is capable of producing 13 of the 22 amino acids, while the other 9, called 'Essential Amino Acids,' must be supplied by food sources." }, { "type": "short_answer", "question": "Name the six essential groups of nutrients broadly classified in the document.", "correct_answer": "Carbohydrates, Proteins, Fats, Vitamins, Minerals, and Water", "explanation": "The document lists Carbohydrates, Proteins, Fats, Vitamins, Minerals, and Water as the six essential groups of nutrients." }, { "type": "multiple_choice", "question": "Which macronutrient is identified as the most concentrated source of energy, providing 9 kcal/g?", "options": [ "Carbohydrates", "Proteins", "Fats", "Vitamins" ], "correct_answer": 2, "explanation": "The document states that 'Fats are the most concentrated source of energy providing 9 kcal/g.'" }, { "type": "true_false", "question": "Water-soluble vitamins, such as B-complex and C, can be stored in the body, while fat-soluble vitamins are easily excreted in urine.", "correct_answer": false, "explanation": "The document states that 'Fat –soluble vitamins can be stored in the body while water soluble vitamins are not and get excreted in urine.'" }, { "type": "short_answer", "question": "What are two primary functions of dietary fiber in the body, as described in the document?", "correct_answer": "Increases gastric motility and aids in digestion; May reduce the risk of developing some diseases like heart disease, diabetes and obesity, and certain cancers.", "explanation": "The document lists 'Increases gastric motility and aids in digestion' and 'May reduce the risk of developing some diseases like heart disease, diabetes and obesity, and certain cancers' as functions of dietary fiber." }, { "type": "multiple_choice", "question": "How is Recommended Dietary Allowance (RDA) defined in the document?", "options": [ "The maximum safe intake level of a nutrient for all individuals.", "The amount of nutrient sufficient for the maintenance of health in nearly all people.", "The average daily nutrient intake level estimated to meet the requirements of half of the healthy individuals in a particular life stage and gender group.", "The minimum amount of a nutrient required to prevent deficiency diseases." ], "correct_answer": 1, "explanation": "RDA is defined as 'the amount of nutrient sufficient for the maintenance of health in nearly all people.'" }, { "type": "true_false", "question": "A diet high in unsaturated fats is associated with a lower level of blood cholesterol and reduces the risk of heart disease.", "correct_answer": true, "explanation": "The document states, 'A diet high in unsaturated fats is associated with a lower level of blood cholesterol and reduces the risk of heart disease.'" }, { "type": "short_answer", "question": "List three physiological functions of food as outlined in the document.", "correct_answer": "Providing energy to carry out voluntary work; Growth or body building; Repair or maintenance of the body cells.", "explanation": "The physiological functions of food include 'Providing energy to carry out voluntary work,' 'Growth or body building,' 'Repair or maintenance of the body cells,' 'Regulation of body processes,' and 'Protective function, increasing one’s resistance to infection.' Any three of these are valid." }, { "type": "multiple_choice", "question": "According to the document, what is the recommended maximum daily intake of salt for adults and children 11 years and over?", "options": [ "Not more than 2g per day", "Not more than 5g per day", "Not more than 10g per day", "Unlimited, as long as blood pressure is normal" ], "correct_answer": 1, "explanation": "The document states, 'It is recommended that adults and children 11 years and over not to have more than 5g of salt per day.'" } ] }
#"""
#render_quiz(json_str)







#already_clicked = st.session_state.get("quiz_key", False)
llm_configuration = LLMConfiguration(st.session_state["selected_llm_provider"], st.session_state["selected_llm_model"])

#generate quiz only if it was not generated before
if quiz_button and "quiz_json_data" not in st.session_state:
    with st.spinner(f"Generating Quiz from {st.session_state['selected_document']}"):
        analysis_document=st.session_state["created_analysis_document"]
        quiz_service=QuizService()
        llm_service:LLMService=LLMFactory.create(llm_configuration)
        document_analysis_orchestrator=DocumentAnalysisOrchestrator(llm_service=llm_service,
                                                                    summary_service=None,
                                                                    quiz_service=quiz_service)
        response=document_analysis_orchestrator.execute(document=analysis_document,action=Action.QUIZ)
        if response:
            #with st.expander("Quiz Generated"):
                #st.markdown(response)
                # 1. Konfiguration der Seite
                #st.set_page_config(page_title="Dokument-Quiz", page_icon="❓", layout="centered")

                #st.title("🧪 Interaktives Quiz")
            st.write("Test your knowledge. Here are the questions from this document. Check your answer by clicking the Link/Expander below the question.")
            #render_quiz(response)
            st.session_state["quiz_json_data"] = json.loads(response)

            # 3. Force Streamlit to redraw the script immediately!
            st.rerun()



# --- ALWAYS CALL THE RENDER METHOD ---
# Every time Streamlit redraws the page when checking an answer, this line runs
# and pulls the questions straight out of st.session_state!
if "quiz_json_data" in st.session_state:
    render_scorable_quiz()